"""
Document ingestion: safe upload, text extraction, chunking.

Location: app/ingestion/documents.py
"""

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

_UNSAFE_CHARS = re.compile(r"[^A-Za-z0-9._-]+")
_REPEATED_DASHES = re.compile(r"-{2,}")


def sanitize_filename(raw: str) -> str:
    """
    Turn an attacker-controlled multipart filename into a safe basename.

    A crafted name like `../../etc/cron.d/x` or `C:\\Windows\\system32\\x`
    must never reach the filesystem, and the extension must stay recognisable.
    """
    name = unicodedata.normalize("NFKD", raw or "").strip()
    name = Path(name.replace("\\", "/")).name
    stem, ext = Path(name).stem, Path(name).suffix.lower()
    stem = _UNSAFE_CHARS.sub("-", stem).strip("-.")[:80] or "document"
    stem = _REPEATED_DASHES.sub("-", stem)
    ext = _UNSAFE_CHARS.sub("", ext)[:10]
    return f"{stem}{ext}"


@dataclass
class ExtractedDocument:
    file_name: str
    file_path: str
    file_type: str
    content: str
    sha256: str

    @property
    def content_length(self) -> int:
        return len(self.content)


def sha256_of_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def save_upload(content: bytes, original_filename: str, dest_dir: Path) -> Path:
    if not content:
        raise ValueError("Empty upload")
    safe_name = sanitize_filename(original_filename)
    if Path(safe_name).suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {Path(safe_name).suffix!r}")
    dest_dir.mkdir(parents=True, exist_ok=True)
    target = dest_dir / safe_name
    target.write_bytes(content)
    return target


def extract_text(path: Path) -> ExtractedDocument:
    ext = path.suffix.lower()
    if ext == ".pdf":
        content = _extract_pdf(path)
    elif ext == ".docx":
        content = _extract_docx(path)
    elif ext in {".txt", ".md"}:
        content = _extract_text(path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

    return ExtractedDocument(
        file_name=path.name,
        file_path=str(path),
        file_type=ext,
        content=content,
        sha256=sha256_of_bytes(path.read_bytes()),
    )


def _extract_pdf(path: Path) -> str:
    parts: List[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"--- Page {page_number} ---\n{text}")
    except Exception:
        try:
            from pypdf import PdfReader

            with open(path, "rb") as handle:
                reader = PdfReader(handle)
                for page_number, page in enumerate(reader.pages, start=1):
                    text = page.extract_text() or ""
                    if text.strip():
                        parts.append(f"--- Page {page_number} ---\n{text}")
        except Exception as exc:
            raise RuntimeError(f"Could not read PDF: {exc}") from exc
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def chunk_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 160,
    metadata: Optional[Dict[str, Any]] = None,
) -> List[Dict[str, Any]]:
    """
    Paragraph-aware chunking with a character overlap so a sentence that
    straddles a boundary is still retrievable from either side.
    """
    if not text or not text.strip():
        return []

    metadata = metadata or {}
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    chunks: List[Dict[str, Any]] = []
    current = ""

    def flush() -> None:
        if not current.strip():
            return
        chunk_meta = dict(metadata)
        chunk_meta["chunk_index"] = len(chunks)
        chunk_meta["chunk_length"] = len(current.strip())
        chunks.append({"text": current.strip(), "metadata": chunk_meta})

    for paragraph in paragraphs:
        if len(paragraph) > chunk_size:
            flush()
            current = ""
            start = 0
            while start < len(paragraph):
                window = paragraph[start : start + chunk_size]
                chunk_meta = dict(metadata)
                chunk_meta["chunk_index"] = len(chunks)
                chunk_meta["chunk_length"] = len(window)
                chunks.append({"text": window.strip(), "metadata": chunk_meta})
                start += max(1, chunk_size - chunk_overlap)
            continue

        if current and len(current) + len(paragraph) + 2 > chunk_size:
            tail = current[-chunk_overlap:] if chunk_overlap else ""
            flush()
            current = (tail + "\n\n" + paragraph).strip() if tail else paragraph
        else:
            current = f"{current}\n\n{paragraph}" if current else paragraph

    flush()
    return chunks


def chunk_document(doc: ExtractedDocument, chunk_size: int, chunk_overlap: int) -> List[Dict[str, Any]]:
    return chunk_text(
        doc.content,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        metadata={"source": doc.file_name, "file_path": doc.file_path, "file_type": doc.file_type},
    )
